#Wilson Lobo, Greivin Mora, Camila Meneses, y Sara Marcano
import os
import csv
from colorama import init, Fore
from procesamiento import limpiar_y_filtrar_lecturas, calcular_metricas_produccion

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

init(autoreset=True)

def main():
    directorio_salida = "Datos Planta"
    os.makedirs(directorio_salida, exist_ok=True)
    
    log_crudo = ["105.5", "98.2", "NaN", "error_sensor", "110.0", "inf", "102.1", "95.4"]
    print(Fore.CYAN + "Procesando log crudo de producción...")
    
    valores_validos, incidencias = limpiar_y_filtrar_lecturas(log_crudo)
    
    ruta_bitacora = os.path.join(directorio_salida, "bitacora_alertas.txt")
    with open(ruta_bitacora, "a", encoding="utf-8") as archivo_bitacora:
        for incidencia in incidencias:
            archivo_bitacora.write(f"[ALERTA] {incidencia}\n")
    print(Fore.YELLOW + f"Bitácora de alertas actualizada en: {ruta_bitacora}")
    
    v_min, v_max, v_prom = calcular_metricas_produccion(valores_validos)
    print(Fore.GREEN + f"Métricas calculadas -> Mín: {v_min}, Máx: {v_max}, Prom: {v_prom:.2f}")
    
    ruta_csv = os.path.join(directorio_salida, "reporte_metricas.csv")
    with open(ruta_csv, "w", newline="", encoding="utf-8") as archivo_csv:
        escritor = csv.writer(archivo_csv, delimiter=";")
        escritor.writerow(["Metrica", "Valor"])
        escritor.writerow(["Minimo", v_min])
        escritor.writerow(["Maximo", v_max])
        escritor.writerow(["Promedio", v_prom])
        
    print(Fore.GREEN + f"Reporte CSV generado exitosamente en: {ruta_csv}")

    wb = Workbook() 
    ws = wb.active
    ws.title = "Producción Mensual" 

    ws["A1"] = "Código Componente" 
    ws["B1"] = "Cantidad Producida"
    
    fuente_negrita = Font(bold=True)
    relleno_azul_marino = PatternFill(start_color="000080", end_color="000080", fill_type="solid") 
    
    ws["A1"].font = fuente_negrita
    ws["A1"].fill = relleno_azul_marino
    ws["B1"].font = fuente_negrita
    ws["B1"].fill = relleno_azul_marino

    datos_simulados = [
        ("COMP-A01", 500),
        ("COMP-B02", 350),
        ("COMP-C03", 820)
    ]
    for registro in datos_simulados:
        ws.append(registro)

    nombre_excel = "Reporte_Calidad.xlsx"
    wb.save(nombre_excel)

    wb_auditoria = load_workbook(nombre_excel) 
    ws_auditoria = wb_auditoria.active
    
    ws_auditoria["B2"] = 485 
    
    wb_auditoria.save(nombre_excel) 
    print(Fore.CYAN + f"Plan Maestro Binario generado y auditado en: {nombre_excel}")

if __name__ == "__main__":
    main()